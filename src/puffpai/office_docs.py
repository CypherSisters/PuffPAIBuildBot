"""Office documents the brain makes and edits, on this computer only.

* ``make_document(kind, title, content)``: a Word (.docx), Excel (.xlsx) or
  PowerPoint (.pptx) file in ``<data folder>/My files/Made by <bot name>/``.
  The answer gives its path, the line "Open it in Office from your files"
  and the one-click open (``POST /files/open``).
* ``edit_document(path, changes)``: simple edits (add paragraphs, rows or
  slides; replace text). A file the bot made (``Made by <bot>``) is changed
  at once; any other file in her data folder only after her yes (one file
  per question); a file in a folder she linked (indexed in place, never
  changed) only after her yes to :data:`LINKED_LINE`, and a copy of the
  original is kept next to it first ("<name> (before BuildBot).<ext>").
  Never anywhere else.

Nothing leaves the computer. New and changed files are read by the data
folder's watcher (:mod:`puffpai.knowledge.my_data`) like any file she adds.
Uses python-docx, openpyxl and python-pptx (requirements.txt).

Copyright (c) 2025 Cypher Sisters. Apache-2.0 License.
"""
from __future__ import annotations

import os
import re
from typing import Any

from puffpai.logging_config import get_logger

logger = get_logger(__name__)

KINDS = {"word": ".docx", "excel": ".xlsx", "powerpoint": ".pptx"}
OPEN_LINE = "Open it in Office from your files"
CONTENT_MAX = 200_000
CHANGES_MAX = 50
LINKED_LINE = ("This file is in a folder you linked. BuildBot normally never changes "
               "those. Change it anyway?")
BEFORE = " (before BuildBot)"


def _bot_name() -> str:
    try:
        from puffpai import bot_identity
        return str(bot_identity.bot_name("") or "").strip()[:60] or "BuildBot"
    except Exception:  # noqa: BLE001
        return "BuildBot"


def _safe(name: Any, fallback: str) -> str:
    text = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', " ", str(name or ""))
    text = " ".join(text.split()).strip(" .")[:80]
    return text or fallback


def made_folder() -> str:
    """``<data folder>/My files/Made by <bot name>`` (made when missing)."""
    from puffpai.knowledge import my_data
    folder = os.path.join(my_data.my_files_dir(create=True), f"Made by {_safe(_bot_name(), 'BuildBot')}")
    os.makedirs(folder, exist_ok=True)
    return folder


def allowed_roots() -> list[str]:
    """Her data folder (My files is inside it) and the folders she linked."""
    from puffpai.knowledge import my_data
    roots = [my_data.data_dir()] + [str(r["path"]) for r in my_data.linked_folders()]
    return [os.path.normcase(os.path.realpath(r)) for r in roots if r]


def allowed_file(path: Any) -> tuple[str, str]:
    """``(real path, "")`` for an existing file inside :func:`allowed_roots`
    (never the vault or BuildBot's own folders), else ``("", plain reason)``."""
    from puffpai.knowledge import my_data
    raw = str(path or "").strip().strip('"')
    if not raw or "\x00" in raw:
        return "", "Say which file."
    real = os.path.realpath(os.path.abspath(os.path.expanduser(raw)))
    if not os.path.isfile(real):
        return "", "I could not find that file."
    canon = os.path.normcase(real)
    inside = False
    for root in allowed_roots():
        try:
            inside = inside or os.path.commonpath([canon, root]) == root
        except ValueError:
            pass
    if not inside or my_data.refusal(os.path.dirname(real)):
        return "", "I can only change files in your data folder or a folder you linked."
    return real, ""


def _inside(canon: str, root: str) -> bool:
    try:
        return os.path.commonpath([canon, root]) == root
    except ValueError:
        return False


def where(real: str) -> str:
    """Where an allowed file is: "made" (the bot's own Made by folder),
    "linked" (a folder she linked) or "mine" (the rest of her data folder)."""
    from puffpai.knowledge import my_data
    canon = os.path.normcase(real)
    made = os.path.normcase(os.path.realpath(os.path.join(
        my_data.my_files_dir(), f"Made by {_safe(_bot_name(), 'BuildBot')}")))
    if _inside(canon, made):
        return "made"
    data = os.path.normcase(os.path.realpath(my_data.data_dir()))
    if _inside(canon, data):
        return "mine"
    return "linked"


def keep_original(real: str) -> str:
    """Copy *real* to "<name> (before BuildBot)<ext>" beside it (a new name
    when one is there already); returns the copy's path."""
    import shutil
    stem, ext = os.path.splitext(os.path.basename(real))
    copy = _unique(os.path.dirname(real), stem + BEFORE, ext)
    shutil.copy2(real, copy)
    return copy


def _unique(folder: str, stem: str, ext: str) -> str:
    path = os.path.join(folder, stem + ext)
    n = 2
    while os.path.exists(path):
        path = os.path.join(folder, f"{stem} ({n}){ext}")
        n += 1
    return path


def _lines(content: Any) -> list[str]:
    if isinstance(content, list):
        return [str(x) if not isinstance(x, list) else ", ".join(map(str, x)) for x in content]
    return str(content or "")[:CONTENT_MAX].splitlines()


def _rows(content: Any) -> list[list[Any]]:
    """Rows for a sheet: a list of lists as given, else one row per line,
    split on tabs, then on commas."""
    if isinstance(content, list):
        return [list(r) if isinstance(r, (list, tuple)) else [r] for r in content]
    out = []
    for line in str(content or "")[:CONTENT_MAX].splitlines():
        if not line.strip():
            continue
        cells = line.split("\t") if "\t" in line else line.split(",")
        out.append([_cell(c.strip()) for c in cells])
    return out


def _cell(text: str) -> Any:
    try:
        return int(text) if re.fullmatch(r"-?\d{1,15}", text) else float(text) \
            if re.fullmatch(r"-?\d+\.\d+", text) else text
    except ValueError:
        return text


def _slides(title: str, content: Any) -> list[tuple[str, list[str]]]:
    """Slides from text: a line "---" or a "# heading" starts a new slide;
    a slide's first line is its title, the rest its points."""
    slides: list[tuple[str, list[str]]] = []
    cur: list[str] = []

    def flush() -> None:
        if cur:
            slides.append((cur[0].lstrip("# ").strip(), [c.lstrip("-* ").strip() for c in cur[1:]]))
    for line in _lines(content):
        s = line.strip()
        if s == "---" or (s.startswith("#") and cur):
            flush()
            cur = [s] if s != "---" else []
            continue
        if s:
            cur.append(s)
    flush()
    return slides or [(title, [])]


# ── writing each kind ───────────────────────────────────────────────────────

def _word_add(doc: Any, lines: list[str]) -> None:
    para: list[str] = []

    def flush() -> None:
        if para:
            doc.add_paragraph(" ".join(para))
            para.clear()
    for line in lines:
        s = line.strip()
        if not s:
            flush()
        elif s.startswith("#"):
            flush()
            level = min(len(s) - len(s.lstrip("#")), 3)
            doc.add_heading(s.lstrip("# ").strip(), level=level)
        elif s.startswith(("- ", "* ")):
            flush()
            doc.add_paragraph(s[2:].strip(), style="List Bullet")
        else:
            para.append(s)
    flush()


def _pptx_add(prs: Any, slides: list[tuple[str, list[str]]]) -> None:
    for head, points in slides:
        layout = prs.slide_layouts[1 if points else 0]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = head
        if points and len(slide.placeholders) > 1:
            body = slide.placeholders[1].text_frame
            body.text = points[0]
            for p in points[1:]:
                body.add_paragraph().text = p


def make_document(kind: str, title: str, content: Any) -> dict:
    kind = str(kind or "").strip().lower()
    if kind not in KINDS:
        return {"ok": False, "said": "Say word, excel or powerpoint."}
    title = " ".join(str(title or "").split())[:120] or "Untitled"
    path = _unique(made_folder(), _safe(title, "Untitled"), KINDS[kind])
    try:
        if kind == "word":
            import docx
            doc = docx.Document()
            doc.add_heading(title, level=0)
            _word_add(doc, _lines(content))
            doc.save(path)
        elif kind == "excel":
            import openpyxl
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = _safe(title, "Sheet")[:31]
            for row in _rows(content):
                ws.append(row)
            wb.save(path)
        else:
            import pptx
            prs = pptx.Presentation()
            slides = _slides(title, content)
            if slides[0][0] != title:
                slides = [(title, [])] + slides
            _pptx_add(prs, slides)
            prs.save(path)
    except ImportError:
        return {"ok": False, "said": f"Making {kind} files needs its Python package; "
                                     "run the installer again to add it."}
    _poke()
    return {"ok": True, "path": path, "name": os.path.basename(path),
            "said": f"Made {os.path.basename(path)} in your files. {OPEN_LINE}.",
            "open": {"method": "POST", "url": "/files/open", "body": {"path": path}}}


def edit_document(path: str, changes: Any, said_yes: bool = False) -> dict:
    """``changes``: a list of ``{"append": text}`` (paragraphs, or slides for
    PowerPoint), ``{"rows": [[...], ...]}`` (Excel) or ``{"replace": old,
    "with": new}``. Only a file the bot made changes without *said_yes*."""
    real, why = allowed_file(path)
    if why:
        return {"ok": False, "said": why}
    ext = os.path.splitext(real)[1].lower()
    if ext not in KINDS.values():
        return {"ok": False, "said": "I can change Word, Excel and PowerPoint files only."}
    place = where(real)
    if place != "made" and not said_yes:
        return {"ok": False, "said": "That file changes only after the owner says yes."}
    todo = [c for c in (changes if isinstance(changes, list) else [changes])
            if isinstance(c, dict)][:CHANGES_MAX]
    if not todo:
        return {"ok": False, "said": "Say what to change."}
    done = 0
    kept = ""
    if place == "linked":
        try:
            kept = keep_original(real)
        except OSError:
            return {"ok": False, "said": "I could not keep a copy of the original, so it "
                                         "was not changed."}
    try:
        if ext == ".docx":
            import docx
            doc = docx.Document(real)
            for c in todo:
                if c.get("replace"):
                    old, new = str(c["replace"]), str(c.get("with") or "")
                    for p in _all_paragraphs(doc):
                        if old in p.text:
                            p.text = p.text.replace(old, new)
                            done += 1
                if c.get("append"):
                    _word_add(doc, _lines(c["append"]))
                    done += 1
            doc.save(real)
        elif ext == ".xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(real)
            ws = wb.active
            for c in todo:
                if c.get("replace"):
                    old, new = str(c["replace"]), str(c.get("with") or "")
                    for sheet in wb.worksheets:
                        for row in sheet.iter_rows():
                            for cell in row:
                                if isinstance(cell.value, str) and old in cell.value:
                                    cell.value = cell.value.replace(old, new)
                                    done += 1
                for key in ("rows", "append"):
                    if c.get(key):
                        for r in _rows(c[key]):
                            ws.append(r)
                            done += 1
            wb.save(real)
        else:
            import pptx
            prs = pptx.Presentation(real)
            for c in todo:
                if c.get("replace"):
                    old, new = str(c["replace"]), str(c.get("with") or "")
                    for slide in prs.slides:
                        for shape in slide.shapes:
                            if not shape.has_text_frame:
                                continue
                            for p in shape.text_frame.paragraphs:
                                for run in p.runs:
                                    if old in run.text:
                                        run.text = run.text.replace(old, new)
                                        done += 1
                if c.get("append"):
                    slides = _slides("", c["append"])
                    _pptx_add(prs, [s for s in slides if s[0] or s[1]])
                    done += 1
            prs.save(real)
    except ImportError:
        return {"ok": False, "said": "Changing that file needs its Python package; run the "
                                     "installer again to add it."}
    except Exception as exc:  # noqa: BLE001 - a file it cannot read
        logger.info("office_docs: %s not changed (%s)", os.path.basename(real), type(exc).__name__)
        return {"ok": False, "said": "I could not open that file to change it."}
    _poke()
    if not done:
        return {"ok": True, "path": real, "said": "Nothing in it matched; it is unchanged."}
    line = f"Changed {os.path.basename(real)}."
    if kept:
        line += f" The original is kept as {os.path.basename(kept)}."
    return {"ok": True, "path": real, "said": f"{line} {OPEN_LINE}.", "original": kept,
            "open": {"method": "POST", "url": "/files/open", "body": {"path": real}}}


def _all_paragraphs(doc: Any) -> list:
    out = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                out += list(cell.paragraphs)
    return out


def _poke() -> None:
    try:
        from puffpai.knowledge import my_data
        my_data.poke()
    except Exception:  # noqa: BLE001 - the next scan finds it anyway
        pass


# ── brain tools ─────────────────────────────────────────────────────────────

def _make_tool(args: dict, ctx: dict) -> dict:
    return make_document(args.get("kind"), args.get("title"), args.get("content"))


def _edit_tool(args: dict, ctx: dict) -> dict:
    return edit_document(args.get("path"), args.get("changes"),
                         said_yes=ctx.get("said_yes") is True)


def edit_question(args: dict, ctx: dict) -> Any:
    """Her yes before a change to any file the bot did not make (one file per
    question); None for the bot's own files and for a file that may not be
    changed at all (the tool says why)."""
    real, why = allowed_file(args.get("path"))
    if why or where(real) == "made":
        return None
    name = os.path.basename(real)
    if where(real) == "linked":
        return {"title": LINKED_LINE,
                "detail": f"{name}, in {os.path.dirname(real)}. A copy of the original is "
                          f"kept next to it first."}
    return {"title": f"May I change {name}?", "detail": f"It is in your files: {real}"}


def register_tools(when: Any = None) -> None:
    from puffpai import brain_tools
    brain_tools.register(
        "make_document",
        "Make a Word, Excel or PowerPoint file on this computer, in the owner's files. "
        "Word: plain text, '# ' headings, '- ' points. Excel: one row per line, cells "
        "split by commas or tabs. PowerPoint: '---' or '# ' starts a slide; a slide's "
        "first line is its title, the rest its points.",
        {"type": "object", "properties": {
            "kind": {"type": "string", "enum": list(KINDS)},
            "title": {"type": "string", "description": "The file's title (its name too)."},
            "content": {"type": "string", "description": "What goes in it."}},
         "required": ["kind", "title", "content"], "additionalProperties": False},
        _make_tool, progress="Making the file ...", when=when)
    brain_tools.register(
        "edit_document",
        "Change a Word, Excel or PowerPoint file in the owner's files or a folder she "
        "linked: add paragraphs, rows or slides, or replace text.",
        {"type": "object", "properties": {
            "path": {"type": "string", "description": "The file's full path."},
            "changes": {"type": "array", "items": {"type": "object", "properties": {
                "append": {"type": "string", "description": "Text to add at the end."},
                "rows": {"type": "array", "items": {"type": "array", "items": {}},
                         "description": "Excel rows to add."},
                "replace": {"type": "string", "description": "Text to find."},
                "with": {"type": "string", "description": "What replaces it."}},
                "additionalProperties": False}}},
         "required": ["path", "changes"], "additionalProperties": False},
        _edit_tool, needs_yes=edit_question, progress="Changing the file ...", when=when)


__all__ = ["BEFORE", "KINDS", "LINKED_LINE", "OPEN_LINE", "edit_question", "keep_original", "where", "allowed_file", "allowed_roots", "edit_document",
           "made_folder", "make_document", "register_tools"]
